from pathlib import Path

from docx import Document


SOURCE = Path("/Users/yun/lindong/天天体育宝课程服务协议0519_修订版.docx")


REPLACEMENTS = {
    "3.  体育活动存在一定的固有风险，如意外扭伤、擦伤等。对于此类非因甲方过错（例如场地隐患、指导失误、管理缺失）所导致的意外，甲方不承担赔偿责任。":
        "3.  体育活动存在一定的固有风险，如意外扭伤、擦伤等。对于因学员自身原因、学员之间正常身体接触或其他非因甲方过错导致的意外，甲方依法不承担赔偿责任；但如损害系因甲方未尽到安全保障、管理或专业指导义务所致，甲方应依法承担相应责任。",
    "4.  因甲方或其雇员（教练）的故意或重大过失行为，或因其提供的场地、器材不符合安全标准，导致学员发生人身损害的，甲方应依法承担相应的赔偿责任。":
        "4.  因甲方或其工作人员存在故意、过失，或因其提供的场地、器材不符合安全标准，导致学员发生人身损害的，甲方应依法承担相应的赔偿责任。",
    "(1) 乙方因个人原因需为学员请假，应至少提前24小时通过甲方指定渠道（小客服微信）告知。":
        "(1) 乙方因个人原因需为学员请假，应至少提前24小时通过甲方指定渠道（如客服微信或平台消息）告知。",
    "(2) 单个学员请假：由于课程为团体授课，单个学员的缺席不影响整体课程正常进行，甲方将正常消耗该学员的当次课时，不予退费，亦不提供录播或补课。但乙方可将该次课程名额转让给符合课程条件的其他亲友使用，需在请假同时将受让人信息（姓名、年龄、监护人联系方式）提前48小时告知甲方。转让完成后，受让人可正常参课。":
        "(2) 单个学员请假：由于课程为团体授课，单个学员缺席通常不影响整体课程正常进行，甲方可正常消耗该学员的当次课时，不予退费，亦不提供录播。但经甲方同意，乙方可将该次课程名额临时转让给符合课程条件的其他亲友使用，并应至少提前24小时提供受让人信息（姓名、年龄、监护人联系方式）。转让完成后，受让人可正常参课。",
    "(2) 开课后解除：课程正式开始后，若乙方因个人原因坚持解除合同并申请退费，甲方应当予以办理。乙方需承担相应的违约责任，违约金金额为合同总金额的30%。甲方在扣除已实际消耗的课时费（按单次课程原价计算）及前述违约金后，将剩余款项退还乙方。":
        "(2) 开课后解除：课程正式开始后，若乙方因个人原因申请解除合同并退费，甲方应予办理。甲方可在退还剩余款项前扣除已实际消耗的课时费；如双方在订单页面、报名须知或补充约定中已明确约定退费手续费或违约责任的，按该等约定执行；未作明确约定的，甲方不得再额外主张违约金。",
    "2.  肖像权使用：甲方可能对训练场景进行拍摄或录像，用于内部教学复盘或宣传推广。乙方有权随时书面通知甲方撤回授权，甲方应在收到通知后停止使用并删除相关素材。":
        "2.  肖像权使用：甲方可能对训练场景进行拍摄或录像，用于内部教学复盘或宣传推广。乙方有权通过书面方式通知甲方不同意或撤回授权；甲方在收到通知后，应停止对相关素材的后续宣传使用，并在合理范围内删除或替换尚未形成公开传播的相关内容。",
    "1.  协议转让：在课程开始前，乙方可将本协议项下的权利义务转让给符合条件的第三方（需满足年龄、体能等要求），需甲方书面同意。":
        "1.  协议转让：在不影响课程正常组织和安全管理的前提下，乙方可在课程开始前申请将本协议项下剩余课程权益转让给符合条件的第三方（需满足年龄、体能等要求），并应取得甲方书面同意。",
    "b.  乙方解除权：出现以下情形之一的，乙方有权要求解除合同并获得剩余款项：":
        "b.  乙方解除权：出现以下情形之一的，乙方有权要求解除合同。退费标准按照本协议第五条第3款执行；如该条已有明确约定的，优先适用该条：",
}


def normalize(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.text = new_text
        return

    first_run = paragraph.runs[0]
    first_run.text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    doc = Document(str(SOURCE))
    original_texts = [normalize(p.text) for p in doc.paragraphs]
    normalized_replacements = {normalize(k): v for k, v in REPLACEMENTS.items()}
    changed = 0
    for paragraph in doc.paragraphs:
        original = normalize(paragraph.text)
        if original in normalized_replacements:
            replace_paragraph_text(paragraph, normalized_replacements[original])
            changed += 1

    if changed != len(REPLACEMENTS):
        missing = set(normalized_replacements) - set(original_texts)
        print(f"warning: expected {len(REPLACEMENTS)} replacements, applied {changed}")
        if missing:
            for item in sorted(missing):
                print("missing:", item)

    doc.save(str(SOURCE))
    print(f"updated {changed} paragraphs in {SOURCE}")


if __name__ == "__main__":
    main()
